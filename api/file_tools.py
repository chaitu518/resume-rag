"""File-system tools (Milestone 1): discover and read resumes / JDs.

Supports .pdf, .docx, .txt and .md. Each loader returns plain text with
line breaks preserved, because the chunker relies on section headers
sitting on their own line.
"""
from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Document:
    path: str
    text: str
    file_type: str


def list_files(directory: str | os.PathLike, exts: set[str] = SUPPORTED_EXTS) -> list[Path]:
    """Recursively list supported files, sorted for deterministic ingestion."""
    root = Path(directory)
    if not root.exists():
        raise FileNotFoundError(f"Directory not found: {root}")
    return sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in exts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def read_file(path: str | os.PathLike) -> Document:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(p)
    elif ext == ".docx":
        text = _read_docx(p)
    elif ext in (".txt", ".md"):
        text = p.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    # normalise whitespace but keep line structure
    lines = [" ".join(line.split()) for line in text.replace("\r", "\n").split("\n")]
    text = "\n".join(lines)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return Document(path=str(p), text=text.strip(), file_type=ext.lstrip("."))


def load_documents(directory: str | os.PathLike) -> list[Document]:
    docs = []
    for f in list_files(directory):
        try:
            docs.append(read_file(f))
        except Exception as e:  # keep going on one bad file
            print(f"[file_tools] skipped {f}: {e}")
    return docs

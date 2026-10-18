"""Generate a diverse synthetic dataset: 36 resumes (PDF/DOCX/TXT), 6 JDs,
and ground-truth relevance labels for evaluation.

Resumes are synthetic (no real people) but deliberately messy: different
section header wording, section order, file formats, skill synonyms
("ML" vs "Machine Learning", "k8s" vs "Kubernetes") and hard distractors
(Python backend devs vs an ML JD, junior vs senior in the same family).

Run:  python -m api.generate_dataset
"""
from __future__ import annotations

import json
import random
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
REF_YEAR, REF_MONTH = 2026, 10
rng = random.Random(42)

FIRST = ["Aarav", "Priya", "Liam", "Sofia", "Wei", "Fatima", "Mateo", "Hana", "Kwame", "Elena",
         "Rohan", "Grace", "Yusuf", "Mei", "Lucas", "Ananya", "Omar", "Chloe", "Diego", "Sakura",
         "Arjun", "Zara", "Noah", "Ifeoma", "Ravi", "Emma", "Kenji", "Leila", "Samuel", "Divya",
         "Ethan", "Nadia", "Vikram", "Olivia", "Tariq", "Isabel", "Karthik", "Amara", "Jonas", "Meera"]
LAST = ["Sharma", "Iyer", "Chen", "Okafor", "Garcia", "Haddad", "Rossi", "Tanaka", "Mensah", "Petrova",
        "Reddy", "Kim", "Nair", "Novak", "Silva", "Kapoor", "Hassan", "Dubois", "Lopez", "Sato",
        "Menon", "Ali", "Fischer", "Eze", "Patel", "Wilson", "Ito", "Rahimi", "Adeyemi", "Rao",
        "Brooks", "Karim", "Joshi", "Martin", "Aziz", "Costa", "Krishnan", "Obi", "Berg", "Pillai"]
COMPANIES = ["Northwind Analytics", "Bluepeak Systems", "Quantia Labs", "Helix Fintech", "Orbital Retail",
             "Cedar Health Tech", "Lumen Logistics", "Vertex Payments", "Nimbus Cloud", "Aurora Media",
             "Ironclad Insurance", "Pixelcraft Studios", "Sagewell Bank", "Trailhead Mobility", "Kestrel AI"]
UNIS = ["IIT Madras", "University of Toronto", "NIT Warangal", "TU Munich", "University of Lagos",
        "National University of Singapore", "BITS Pilani", "University of Michigan", "Anna University",
        "University of Melbourne", "IIIT Hyderabad", "Georgia Tech"]

FAMILIES = {
    "ml_engineer": {
        "titles": ["Machine Learning Engineer", "Senior ML Engineer", "Data Scientist", "Applied Scientist"],
        "core": ["Python", "Machine Learning", "PyTorch", "scikit-learn", "Pandas", "NLP"],
        "extra": ["TensorFlow", "Deep Learning", "MLOps", "AWS", "Docker", "LLMs", "RAG", "Spark", "SQL", "Computer Vision"],
        "bullets": [
            "Built and deployed {s1} models in {s0} serving 2M+ predictions per day with p95 latency under 80ms",
            "Trained {s2} classifiers for fraud detection, lifting recall by 18% over the rules baseline",
            "Designed feature pipelines with {s4} and {s3}, cutting training data prep time from days to hours",
            "Fine-tuned transformer models for {s5} tasks such as entity extraction and intent classification",
            "Set up experiment tracking and model registry (MLflow) so the team could reproduce every run",
            "Led A/B tests for a recommendation model that raised click-through rate by 7%",
            "Productionised a retrieval-augmented question answering service over 40k internal documents",
            "Mentored two junior data scientists and reviewed modelling code for statistical rigour",
        ],
        "degree": ["M.Tech in Artificial Intelligence", "M.S. in Computer Science", "Ph.D. in Machine Learning",
                   "B.Tech in Computer Science"],
        "projects": ["Resume screening classifier using sentence embeddings", "Churn prediction with gradient boosting",
                     "Medical image segmentation with U-Net"],
        "certs": ["AWS Certified Machine Learning - Specialty", "DeepLearning.AI Deep Learning Specialization"],
    },
    "java_backend": {
        "titles": ["Backend Engineer", "Senior Java Developer", "Software Engineer - Backend", "SDE 2"],
        "core": ["Java", "Spring Boot", "Microservices", "REST APIs", "MySQL", "Kafka"],
        "extra": ["Redis", "Docker", "Kubernetes", "AWS", "Hibernate", "PostgreSQL", "System Design", "JUnit", "MongoDB"],
        "bullets": [
            "Developed {s1} microservices in {s0} 17 handling 15k requests per second for the payments platform",
            "Migrated a monolith to event-driven {s2} using {s5}, reducing order processing delay by 60%",
            "Designed {s3} consumed by web and mobile clients, with contract tests and versioning",
            "Tuned {s4} queries and added Redis caching, cutting p99 latency from 900ms to 120ms",
            "Implemented idempotent consumers and dead-letter queues for reliable event handling",
            "Owned on-call for three services and wrote runbooks that halved mean time to recovery",
            "Introduced unit and integration testing standards, raising coverage from 45% to 82%",
        ],
        "degree": ["B.Tech in Computer Science", "B.E. in Information Technology", "M.S. in Software Engineering", "MCA"],
        "projects": ["Ticket booking system with seat locking", "Rate limiter library using token bucket",
                     "Library management system LLD"],
        "certs": ["Oracle Certified Professional Java SE 17", "AWS Certified Developer - Associate"],
    },
    "python_backend": {
        "titles": ["Python Developer", "Backend Engineer (Python)", "Software Engineer"],
        "core": ["Python", "Django", "FastAPI", "PostgreSQL", "REST APIs", "Docker"],
        "extra": ["Redis", "Celery", "AWS", "Flask", "RabbitMQ", "Kubernetes", "Pytest", "GraphQL"],
        "bullets": [
            "Built {s1} and {s2} services in {s0} powering the customer onboarding flow",
            "Designed {s4} for partner integrations with OAuth2 and rate limiting",
            "Optimised {s3} schemas and indexes, reducing report generation time by 70%",
            "Containerised services with {s5} and automated deployments through CI pipelines",
            "Wrote async background workers for invoice generation processing 500k jobs per day",
            "Added structured logging and tracing, making production debugging much faster",
        ],
        "degree": ["B.Tech in Computer Science", "B.Sc in Computer Science", "BCA"],
        "projects": ["URL shortener with analytics", "Expense tracker API"],
        "certs": ["Python Institute PCAP"],
    },
    "data_engineer": {
        "titles": ["Data Engineer", "Senior Data Engineer", "Big Data Engineer"],
        "core": ["Python", "Spark", "Airflow", "SQL", "ETL", "Snowflake"],
        "extra": ["Kafka", "AWS", "dbt", "Hadoop", "Scala", "GCP", "Docker", "Terraform"],
        "bullets": [
            "Built batch and streaming {s4} pipelines with {s1} processing 3TB of events per day",
            "Orchestrated 200+ DAGs in {s2} with SLAs, retries and data quality checks",
            "Modelled warehouse layers in {s5} using dbt, giving analysts trusted, documented tables",
            "Rewrote legacy Hive jobs in {s1} (PySpark), cutting compute cost by 35%",
            "Implemented CDC ingestion from operational databases using Debezium",
            "Partnered with ML engineers to deliver feature tables for model training",
        ],
        "degree": ["B.Tech in Computer Science", "M.S. in Data Science", "B.E. in Electronics"],
        "projects": ["Real-time clickstream pipeline", "Data lake on S3 with Iceberg tables"],
        "certs": ["Databricks Certified Data Engineer Associate", "Google Professional Data Engineer"],
    },
    "devops": {
        "titles": ["DevOps Engineer", "Site Reliability Engineer", "Platform Engineer", "Cloud Engineer"],
        "core": ["Kubernetes", "Docker", "Terraform", "AWS", "CI/CD", "Linux"],
        "extra": ["Prometheus", "Ansible", "Python", "Bash", "Azure", "GCP", "Go", "Kafka"],
        "bullets": [
            "Managed 40+ {s0} clusters (EKS) running 600 services with zero-downtime upgrades",
            "Wrote {s2} modules for VPCs, IAM and databases, making all {s3} infra reproducible",
            "Built {s4} pipelines with GitHub Actions and ArgoCD, cutting release time from 2 days to 30 minutes",
            "Introduced SLOs and alerting with Prometheus and Grafana, reducing pager noise by 50%",
            "Hardened {s5} base images and container builds with {s1}, passing SOC2 audit",
            "Led incident reviews and capacity planning for peak traffic events",
        ],
        "degree": ["B.Tech in Information Technology", "B.E. in Computer Science", "Diploma in Network Engineering"],
        "projects": ["Self-service preview environments on Kubernetes", "Cost dashboard for cloud spend"],
        "certs": ["Certified Kubernetes Administrator (CKA)", "HashiCorp Terraform Associate",
                  "AWS Certified Solutions Architect - Associate"],
    },
    "frontend": {
        "titles": ["Frontend Engineer", "Senior React Developer", "UI Engineer", "Full Stack Developer"],
        "core": ["React", "TypeScript", "JavaScript", "HTML/CSS", "Redux", "Next.js"],
        "extra": ["Node.js", "GraphQL", "Cypress", "Vue.js", "Angular", "Docker", "AWS"],
        "bullets": [
            "Built a component library in {s0} and {s1} used by 6 product teams",
            "Migrated a legacy SPA to {s5} with server-side rendering, improving LCP by 40%",
            "Managed complex client state with {s4} and React Query for a trading dashboard",
            "Implemented accessible, responsive layouts with {s3} meeting WCAG 2.1 AA",
            "Wrote end-to-end tests with Cypress and raised release confidence",
            "Worked with designers to ship a new onboarding flow that increased activation by 12%",
        ],
        "degree": ["B.Tech in Computer Science", "B.Sc in Information Technology", "B.Des in Interaction Design"],
        "projects": ["Real-time collaborative whiteboard", "Personal portfolio with Next.js"],
        "certs": ["Meta Front-End Developer Certificate"],
    },
    "mobile": {
        "titles": ["Android Developer", "iOS Engineer", "Mobile Engineer"],
        "core": ["Kotlin", "Android", "Swift", "iOS", "Flutter", "REST APIs"],
        "extra": ["React Native", "Java", "CI/CD", "GraphQL"],
        "bullets": [
            "Shipped {s1} features in {s0} for an app with 5M+ installs and 4.6 store rating",
            "Built cross-platform screens with {s4}, sharing 70% of code between {s1} and {s3}",
            "Reduced app start time by 35% through lazy loading and image caching",
            "Integrated {s5} with offline sync and conflict resolution",
        ],
        "degree": ["B.Tech in Computer Science", "B.E. in Electronics and Communication"],
        "projects": ["Habit tracker app", "Offline-first notes app"],
        "certs": ["Google Associate Android Developer"],
    },
    "qa": {
        "titles": ["QA Automation Engineer", "SDET", "Test Engineer"],
        "core": ["Selenium", "Test Automation", "Java", "Postman", "JUnit", "Cypress"],
        "extra": ["Python", "Pytest", "CI/CD", "SQL", "Agile"],
        "bullets": [
            "Built a {s1} framework with {s0} and {s2} covering 1,200 regression cases",
            "Automated API checks with {s3}, catching contract breaks before release",
            "Integrated test suites into CI so every pull request ran smoke tests in 8 minutes",
            "Reduced flaky tests by 80% through better waits and test data isolation",
        ],
        "degree": ["B.Tech in Information Technology", "B.Sc in Computer Science"],
        "projects": ["Visual regression testing tool"],
        "certs": ["ISTQB Certified Tester Foundation Level"],
    },
    "data_analyst": {
        "titles": ["Data Analyst", "Business Intelligence Analyst", "Product Analyst"],
        "core": ["SQL", "Tableau", "Excel", "Power BI", "Statistics", "Python"],
        "extra": ["Pandas", "Snowflake", "R", "GCP"],
        "bullets": [
            "Wrote complex {s0} queries and built {s1} dashboards tracking revenue and retention",
            "Ran A/B tests and {s4} reviews for pricing experiments across 3 markets",
            "Automated weekly reporting in {s2} and {s5}, saving the team 10 hours a week",
            "Built {s3} reports for leadership on funnel conversion and churn drivers",
        ],
        "degree": ["B.Com with Business Analytics", "B.Sc in Statistics", "MBA in Business Analytics",
                   "B.Tech in Computer Science"],
        "projects": ["Customer segmentation with k-means", "Sales forecasting dashboard"],
        "certs": ["Google Data Analytics Professional Certificate", "Tableau Desktop Specialist"],
    },
}

# skills every resume in a family lists (the "primary" stack for that role)
MUST_LIST = {"ml_engineer": ["Python"], "java_backend": ["Java", "Spring Boot"], "data_engineer": ["Spark"],
             "devops": ["Kubernetes", "Terraform"], "frontend": ["React", "TypeScript"], "data_analyst": ["SQL"],
             "python_backend": ["Python"], "qa": ["Selenium"], "mobile": []}

# (family, years) — 36 resumes, deliberate spread of seniority per family
ROSTER = [
    ("ml_engineer", 8), ("ml_engineer", 6), ("ml_engineer", 5.5), ("ml_engineer", 3), ("ml_engineer", 1.5),
    ("java_backend", 9), ("java_backend", 6), ("java_backend", 4.5), ("java_backend", 3.5), ("java_backend", 2),
    ("python_backend", 7), ("python_backend", 5), ("python_backend", 2.5),
    ("data_engineer", 7), ("data_engineer", 5), ("data_engineer", 3.5), ("data_engineer", 2),
    ("devops", 10), ("devops", 6), ("devops", 4), ("devops", 2.5),
    ("frontend", 8), ("frontend", 5), ("frontend", 3), ("frontend", 1.5),
    ("mobile", 6), ("mobile", 3), ("mobile", 2),
    ("qa", 7), ("qa", 4), ("qa", 1.5),
    ("data_analyst", 6), ("data_analyst", 4), ("data_analyst", 2.5), ("data_analyst", 1),
    ("ml_engineer", 4),
]

HEADER_VARIANTS = {
    "summary": ["SUMMARY", "Professional Summary", "Profile", "About Me"],
    "experience": ["EXPERIENCE", "Work Experience", "Professional Experience", "Employment History"],
    "education": ["EDUCATION", "Education", "Academic Background"],
    "skills": ["SKILLS", "Technical Skills", "Core Skills", "Tech Stack"],
    "projects": ["PROJECTS", "Key Projects", "Personal Projects"],
    "certifications": ["CERTIFICATIONS", "Certifications", "Licenses & Certifications"],
}
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _fmt(idx: int) -> str:  # month index -> "Mar 2021"
    y, m = divmod(idx - 1, 12)
    return f"{MON[m]} {y}"


def build_jobs(years: float, fam: dict, skills: list[str], extras: list[str]):
    total = int(round(years * 12))
    n_jobs = 1 if years < 3 else (2 if years < 6 else 3)
    cuts = sorted(rng.sample(range(6, total - 5), n_jobs - 1)) if n_jobs > 1 else []
    bounds = [0] + cuts + [total]
    end_idx = REF_YEAR * 12 + REF_MONTH
    jobs = []
    companies = rng.sample(COMPANIES, n_jobs)
    pool = rng.sample(fam["bullets"], len(fam["bullets"]))
    for j in range(n_jobs):  # most recent first
        dur_end = end_idx - bounds[j]
        dur_start = end_idx - bounds[j + 1]
        title = fam["titles"][min(j, len(fam["titles"]) - 1)] if j else rng.choice(fam["titles"])
        fill = {f"s{i}": s for i, s in enumerate(skills)}
        fill.update({f"x{i}": e for i, e in enumerate(extras + ["MLflow", "Debezium"])})
        k = max(2, min(rng.randint(3, 4), len(fam["bullets"]) // n_jobs))
        if len(pool) < k:
            pool.extend(rng.sample(fam["bullets"], len(fam["bullets"])))
        bl = [pool.pop() for _ in range(k)]
        bullets = []
        for b in bl:
            try:
                bullets.append(b.format(**fill))
            except (KeyError, IndexError):
                bullets.append(b)
        jobs.append({
            "title": title, "company": companies[j],
            "start": _fmt(dur_start + 1), "end": "Present" if j == 0 else _fmt(dur_end),
            "bullets": bullets,
        })
    return jobs


def make_resume(i: int, family: str, years: float) -> dict:
    fam = FAMILIES[family]
    name = f"{FIRST[i]} {LAST[(i * 7) % len(LAST)]}"
    must = MUST_LIST.get(family, [])
    rest = [c for c in fam["core"] if c not in must]
    rng.shuffle(rest)
    core = must + rest[: (rng.randint(4, 6) if years > 2 else 4) - len(must)]
    extras = rng.sample(fam["extra"], k=min(len(fam["extra"]), 2 + int(years // 3)))
    # bullets reference skills by position; keep the family's canonical order for readability
    jobs = build_jobs(years, fam, fam["core"], extras)
    # junior candidates rarely list everything
    listed_skills = [s for s in fam["core"] if s in core] + extras
    # synonym noise so keyword search alone isn't enough
    synonyms = {"Machine Learning": "ML", "Kubernetes": "k8s", "JavaScript": "JS (ES6+)",
                "PostgreSQL": "Postgres", "scikit-learn": "sklearn"}
    shown = [synonyms[s] if s in synonyms and rng.random() < 0.4 else s for s in listed_skills]
    degree = rng.choice(fam["degree"])
    grad_year = REF_YEAR - int(years) - rng.randint(0, 2)
    title = jobs[0]["title"]
    yrs_txt = f"{int(years)}+" if years >= 2 else "1+"
    summary = rng.choice([
        f"{title} with {yrs_txt} years of experience building production systems using {', '.join(listed_skills[:3])}.",
        f"Results-driven {title.lower()} ({yrs_txt} yrs) focused on {listed_skills[0]} and {listed_skills[1]}. "
        f"Comfortable owning features end to end and working with cross-functional teams.",
        f"Engineer with {yrs_txt} years in {listed_skills[0]}, {listed_skills[1]} and {listed_skills[2]}, "
        f"passionate about clean design, measurable impact and mentoring.",
    ])
    if years < 2:
        summary = f"Early-career {title.lower()} with about a year of hands-on {listed_skills[0]} work and strong fundamentals."
    sections = {
        "summary": summary,
        "experience": "\n\n".join(
            f"{j['title']} - {j['company']}\n{j['start']} - {j['end']}\n" + "\n".join(f"- {b}" for b in j["bullets"])
            for j in jobs),
        "education": f"{degree}, {rng.choice(UNIS)} ({grad_year})",
        "skills": ", ".join(shown),
        "projects": "\n".join(f"- {p}" for p in rng.sample(fam["projects"], k=min(2, len(fam["projects"])))),
    }
    if years >= 3 and rng.random() < 0.7:
        sections["certifications"] = "\n".join(f"- {c}" for c in rng.sample(fam["certs"], k=1))
    order = ["summary", "experience", "skills", "education", "projects", "certifications"]
    if rng.random() < 0.4:
        order = ["summary", "skills", "experience", "projects", "education", "certifications"]
    email = f"{name.split()[0].lower()}.{name.split()[1].lower()}@example.com"
    phone = f"+1-555-{rng.randint(100, 999)}-{rng.randint(1000, 9999)}"
    return {
        "id": f"{name.lower().replace(' ', '_')}",
        "name": name, "family": family, "years": years, "title": title,
        "header": [name, title, f"{email} | {phone} | linkedin.com/in/{name.split()[0].lower()}"],
        "sections": [(HEADER_VARIANTS[s][rng.randrange(len(HEADER_VARIANTS[s]))], sections[s]) for s in order if s in sections],
        "true_skills": listed_skills, "degree": degree,
    }


def to_text(r: dict) -> str:
    parts = ["\n".join(r["header"])]
    for head, body in r["sections"]:
        parts.append(f"{head}\n{body}")
    return "\n\n".join(parts) + "\n"


def write_pdf(r: dict, path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Heading3"], spaceBefore=8, spaceAfter=2)
    story = [Paragraph(escape(r["header"][0]), ss["Title"]),
             Paragraph(escape(r["header"][1]), ss["Normal"]),
             Paragraph(escape(r["header"][2]), ss["Normal"]), Spacer(1, 6)]
    for head, body in r["sections"]:
        story.append(Paragraph(escape(head), h))
        for line in body.split("\n"):
            if line.strip():
                story.append(Paragraph(escape(line), ss["Normal"]))
            else:
                story.append(Spacer(1, 4))
    SimpleDocTemplate(str(path), pagesize=A4, title=r["name"]).build(story)


def write_docx(r: dict, path: Path):
    import docx
    d = docx.Document()
    d.add_heading(r["header"][0], level=0)
    d.add_paragraph(r["header"][1])
    d.add_paragraph(r["header"][2])
    for head, body in r["sections"]:
        d.add_heading(head, level=2)
        for line in body.split("\n"):
            if line.strip():
                d.add_paragraph(line)
    d.save(str(path))


JOB_DESCRIPTIONS = [
    {"id": "jd01_senior_ml_engineer", "family": "ml_engineer", "min_years": 5, "must_have": ["Python"],
     "text": """Senior Machine Learning Engineer - Kestrel AI
We are looking for a Senior ML Engineer to build and ship models that power our recommendation and NLP products.

Responsibilities:
- Design, train and deploy machine learning models to production
- Build NLP features such as entity extraction and semantic search using transformers and LLMs
- Own the MLOps lifecycle: experiment tracking, model registry, monitoring

Requirements:
- 5+ years Python experience
- Strong background in Machine Learning and Deep Learning (PyTorch or TensorFlow)
- Experience with NLP and retrieval systems (RAG is a plus)
- Familiarity with AWS, Docker and MLOps tooling
"""},
    {"id": "jd02_java_backend", "family": "java_backend", "min_years": 4, "must_have": ["Java", "Spring Boot"],
     "text": """Backend Engineer (Java) - Helix Fintech
Join our payments team building high-throughput, event-driven services.

What you'll do:
- Build Spring Boot microservices and REST APIs used by millions of customers
- Design event-driven flows with Kafka
- Tune MySQL/PostgreSQL and Redis for low latency

Must have:
- 4+ years Java
- Spring Boot

Also expected: Microservices, Kafka, REST APIs.
Nice to have: Kubernetes, AWS, System Design experience
"""},
    {"id": "jd03_data_engineer", "family": "data_engineer", "min_years": 3, "must_have": ["Spark"],
     "text": """Data Engineer - Orbital Retail
We need a data engineer to build reliable batch and streaming pipelines for analytics and ML.

Responsibilities:
- Build ETL/ELT pipelines with Spark and Airflow
- Model warehouse tables in Snowflake with dbt
- Ensure data quality, SLAs and documentation

Requirements:
- 3+ years of data engineering experience
- Strong SQL and Python, hands-on Spark
- Experience with Airflow and a cloud warehouse (Snowflake/BigQuery)
"""},
    {"id": "jd04_devops_sre", "family": "devops", "min_years": 4, "must_have": ["Kubernetes", "Terraform"],
     "text": """Site Reliability / DevOps Engineer - Nimbus Cloud
Help us run a reliable, secure platform on AWS.

You will:
- Operate Kubernetes clusters and container workloads at scale
- Write infrastructure as code with Terraform
- Build CI/CD pipelines and observability (Prometheus/Grafana)

Requirements:
- 4+ years in DevOps/SRE roles
- Kubernetes and Terraform are must-haves
- AWS, Docker, Linux, scripting in Python or Bash
"""},
    {"id": "jd05_frontend_react", "family": "frontend", "min_years": 3, "must_have": ["React", "TypeScript"],
     "text": """Frontend Engineer (React) - Pixelcraft Studios
Build fast, accessible web experiences for our creative tools.

Responsibilities:
- Develop UI features in React and TypeScript
- Own our component library and design system
- Improve performance with Next.js server-side rendering

Requirements:
- 3+ years React
- Strong TypeScript, JavaScript, HTML/CSS
- State management (Redux or similar), testing with Cypress a plus
"""},
    {"id": "jd06_data_analyst", "family": "data_analyst", "min_years": 2, "must_have": ["SQL"],
     "text": """Data Analyst - Sagewell Bank
Turn data into decisions for product and leadership teams.

Responsibilities:
- Write SQL to answer business questions
- Build dashboards in Tableau or Power BI
- Run A/B tests and statistical analysis

Requirements:
- 2+ years as a data or BI analyst
- Advanced SQL and Excel
- Tableau or Power BI, solid statistics
"""},
]


# Short, paraphrased recruiter queries: few exact skill names, tests semantic recall
PARAPHRASED_QUERIES = [
    {"id": "q01", "family": "ml_engineer", "text": "someone who trains neural networks for text understanding and puts models into production"},
    {"id": "q02", "family": "java_backend", "text": "engineer building high-throughput payment services on the JVM with event streaming"},
    {"id": "q03", "family": "data_engineer", "text": "build and orchestrate warehouse pipelines that move terabytes of events every day"},
    {"id": "q04", "family": "devops", "text": "keep our container clusters healthy, write infrastructure as code and own release automation"},
    {"id": "q05", "family": "frontend", "text": "craft accessible, fast web interfaces and a reusable component library"},
    {"id": "q06", "family": "data_analyst", "text": "analyst who builds dashboards for leadership and runs pricing experiments"},
    {"id": "q07", "family": "mobile", "text": "ship features in our phone apps on both app stores with offline sync"},
    {"id": "q08", "family": "qa", "text": "own regression test automation and catch API contract breaks before release"},
]


def relevance(candidate: dict, jd: dict) -> int:
    """Graded relevance: 2 = right family & meets years, 1 = right family but junior, 0 = other."""
    if candidate["family"] != jd["family"]:
        return 0
    return 2 if candidate["years"] >= jd["min_years"] else 1


def main():
    res_dir = DATA / "resumes"
    jd_dir = DATA / "job_descriptions"
    for d in (res_dir, jd_dir):
        d.mkdir(parents=True, exist_ok=True)
        for f in d.iterdir():
            f.unlink()
    resumes, truth = [], {}
    for i, (family, years) in enumerate(ROSTER):
        r = make_resume(i, family, years)
        fmt = ["pdf", "pdf", "docx", "txt"][i % 4]
        path = res_dir / f"{r['id']}.{fmt}"
        if fmt == "pdf":
            write_pdf(r, path)
        elif fmt == "docx":
            write_docx(r, path)
        else:
            path.write_text(to_text(r))
        rel_path = str(path.relative_to(ROOT))
        r["path"] = rel_path
        resumes.append(r)
        truth[rel_path] = {"name": r["name"], "family": family, "years": years,
                           "skills": r["true_skills"], "degree": r["degree"]}
    for jd in JOB_DESCRIPTIONS:
        (jd_dir / f"{jd['id']}.txt").write_text(jd["text"])
    gt = {
        "reference_date": f"{REF_YEAR}-{REF_MONTH:02d}",
        "candidates": truth,
        "job_descriptions": {
            jd["id"]: {
                "path": f"data/job_descriptions/{jd['id']}.txt",
                "family": jd["family"], "min_years": jd["min_years"], "must_have": jd["must_have"],
                "relevance": {r["path"]: relevance(r, jd) for r in resumes if relevance(r, jd) > 0},
            } for jd in JOB_DESCRIPTIONS
        },
    }
    gt["paraphrased_queries"] = {
        q["id"]: {"text": q["text"], "family": q["family"],
                  "relevance": {r["path"]: 1 for r in resumes if r["family"] == q["family"]}}
        for q in PARAPHRASED_QUERIES}
    (DATA / "ground_truth.json").write_text(json.dumps(gt, indent=2))
    fmts = {}
    for r in resumes:
        fmts[r["path"].rsplit(".", 1)[1]] = fmts.get(r["path"].rsplit(".", 1)[1], 0) + 1
    print(f"Wrote {len(resumes)} resumes {fmts} and {len(JOB_DESCRIPTIONS)} job descriptions to {DATA}")


if __name__ == "__main__":
    main()
